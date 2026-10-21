#!/usr/bin/env python3
"""
FeFET endurance figures from the Clarius dual-sweep workbooks (``554*``, ``556*``).

    python plot_fefet_endurance.py                    # every matching workbook
    python plot_fefet_endurance.py "554 3-13#1.xlsx"  # just these
    python plot_fefet_endurance.py --outdir DIR --ic 1e-9

Each workbook holds one sheet per endurance checkpoint, and each sheet is a full
Id-Vg dual sweep (Vg -5.5 -> +5.5 -> -5.5 V, step 0.02 V, Vd = 0.1 V, from the
``Settings`` sheet). Sheet names carry the cycle count:

    ``pristine`` / ``... pr``   -> 0 cycles (as fabricated)
    ``10^5``                    -> 1e5 cycles
    ``556 2-17 5``              -> 1e5 cycles (trailing token read as the exponent)

The mapping is printed at run time -- check it, it is an inference from the names.

Per workbook, ``fefet_endurance_<stem>.png``:
  (a) Id-Vg transfer curves, both sweep branches, shaded by cycle count
  (b) forward / reverse threshold voltage vs cycles
  (c) memory window (Vth_fwd - Vth_rev) vs cycles
  (d) on-current at max Vg vs cycles
  (e) subthreshold swing vs cycles

Panels (a), (b) and (c) are also written on their own, without the panel letter,
as ``fefet_endurance_<stem>_transfer.png``, ``..._vth.png`` and ``..._window.png``
(``--no-panels`` skips them).

Across workbooks, ``fefet_endurance_comparison.png``: memory window in volts and
as a percentage of the pristine window.

Vth is the constant-current crossing at ``IC_A`` (1 nA at Vd = 0.1 V), taken per
branch; the off-current is not plotted because it sits at the SMU noise floor
(~1e-15 A), so an on/off ratio here would be a measure of the instrument.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import Optional

import matplotlib
matplotlib.use("Agg")           # save-only; no display needed (WSL, ssh, CI)

import matplotlib.pyplot as plt          # noqa: E402
import numpy as np                       # noqa: E402
import openpyxl                          # noqa: E402
from matplotlib.colors import Normalize  # noqa: E402

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parent
OUTDIR = ROOT / "figures"
DEFAULT_GLOBS = ("554*.xlsx", "556*.xlsx")

IC_A = 1e-9             # constant-current criterion for Vth, in amperes
SS_ID_MIN = 1e-12       # subthreshold-swing fit window, lower bound [A].
                        # The SMU floor is ~1e-15 A; fitting nearer to it makes
                        # the minimum slope a noise measurement, not an SS.
SS_ID_FRAC = 0.05       # ... upper bound as a fraction of Ion
SS_HALF_WIN = 10        # half-width, in points, of the local slope fit
ION_STEP_WARN = 0.20    # flag an Ion jump larger than this between checkpoints

FIG_SIZE = (18, 10)
FIG_CMP_SIZE = (13, 5.5)
FIG_TRANSFER_SIZE = (9, 6.5)    # panel (a) on its own
FIG_PANEL_SIZE = (8, 6)         # panels (b), (c) on their own
DPI = 300

# Style (aligned with plot_fefet_retention.py / plot_igzo_4t2c_retention.py)
FONT_BASE = 18
FONT_AXIS = 18
FONT_PANEL = 18
LW_CURVE = 2.0
GRID_ALPHA = 0.5
MARKERSIZE = 7

# Cycle count is a magnitude, so it gets one perceptually ordered ramp plus a
# colorbar -- not ten legend entries. The pristine sweep is off the ramp entirely
# and is not drawn in the transfer panel; it sets the reference in the others.
CMAP_CYCLES = "viridis"
C_FWD = "#0072B2"       # forward branch / device 1
C_REV = "#D55E00"       # reverse branch / device 2
C_MW = "#0072B2"
C_ION = "#009E73"
C_SS = "#CC3311"
DEVICE_COLORS = ("#0072B2", "#D55E00", "#009E73", "#CC3311")

PRISTINE_X = 0.03       # where "0 cycles" sits on the log cycle axis


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------
def cycles_from_sheet(name: str) -> Optional[float]:
    """Cycle count encoded in a sheet name, or None if it is not a data sheet."""
    s = name.strip().lower()
    if "pristine" in s or re.search(r"(^|\s)pr$", s):
        return 0.0
    m = re.search(r"10\s*\^\s*(\d+)", s)
    if m:
        return float(10 ** int(m.group(1)))
    m = re.search(r"(?:^|\s)(\d{1,2})\s*$", s)     # trailing token = exponent
    if m:
        return float(10 ** int(m.group(1)))
    return None


def read_sweeps(xlsx_path: Path) -> list:
    """Return [{cycles, sheet, vg, id}] for every Id-Vg sheet, pristine first."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        header = [c.value for c in next(ws.iter_rows(max_row=1), ())]
        if len(header) < 7 or header[0] != "Time" or "GateV" not in header:
            continue                                    # Calc / Settings sheets
        i_gate, i_drain = header.index("GateV"), header.index("DrainI")
        cycles = cycles_from_sheet(ws.title)
        if cycles is None:
            print(f"[warn] {xlsx_path.name}: cannot read a cycle count from "
                  f"sheet {ws.title!r}; skipped")
            continue
        vg, idr = [], []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[i_gate] is None or row[i_drain] is None:
                continue
            try:                                        # cells are stored as text
                vg.append(float(row[i_gate]))
                idr.append(float(row[i_drain]))
            except (TypeError, ValueError):
                continue
        if len(vg) < 10:
            print(f"[warn] {xlsx_path.name}: sheet {ws.title!r} has no usable data")
            continue
        out.append({"cycles": cycles, "sheet": ws.title,
                    "vg": np.asarray(vg), "id": np.asarray(idr)})
    wb.close()
    out.sort(key=lambda d: d["cycles"])
    return out


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------
def split_branches(vg: np.ndarray, idr: np.ndarray) -> tuple:
    """Split a dual sweep at its turning point into (forward, reverse)."""
    turn = int(np.argmax(vg)) + 1
    return (vg[:turn], idr[:turn]), (vg[turn:], idr[turn:])


def vth_constant_current(vg: np.ndarray, idr: np.ndarray, ic: float = IC_A) -> float:
    """Vg where |Id| first crosses ``ic``, interpolated in log10(Id)."""
    a = np.abs(np.asarray(idr, float))
    order = np.argsort(vg)
    v, a = np.asarray(vg, float)[order], a[order]
    above = np.flatnonzero(a >= ic)
    if above.size == 0 or above[0] == 0:
        return float("nan")                 # never crosses, or starts above ic
    i = above[0]
    lo, hi = np.log10(max(a[i - 1], 1e-18)), np.log10(a[i])
    if hi <= lo:
        return float(v[i])
    return float(v[i - 1] + (np.log10(ic) - lo) / (hi - lo) * (v[i] - v[i - 1]))


def subthreshold_swing(vg: np.ndarray, idr: np.ndarray, ion: float) -> float:
    """Minimum dVg/dlog10|Id| [mV/dec] over the subthreshold window."""
    a = np.abs(np.asarray(idr, float))
    order = np.argsort(vg)
    v, a = np.asarray(vg, float)[order], a[order]
    hi = max(SS_ID_FRAC * ion, SS_ID_MIN * 10)
    win = np.flatnonzero((a > SS_ID_MIN) & (a < hi))
    if win.size < 2 * SS_HALF_WIN + 1:
        return float("nan")
    logi = np.log10(a)
    slopes = []
    for k in win[SS_HALF_WIN:-SS_HALF_WIN]:
        sl = slice(k - SS_HALF_WIN, k + SS_HALF_WIN + 1)
        if not np.all(np.isfinite(logi[sl])):
            continue
        d_logi, _ = np.polyfit(v[sl], logi[sl], 1)
        if d_logi > 0:
            slopes.append(1.0 / d_logi)     # V/dec
    return float(np.min(slopes) * 1e3) if slopes else float("nan")


def metrics(sweeps: list, ic: float = IC_A) -> dict:
    """Per-checkpoint Vth (both branches), memory window, Ion and SS."""
    cycles, vth_f, vth_r, ion, ss, sheets = [], [], [], [], [], []
    for s in sweeps:
        (vf, if_), (vr, ir) = split_branches(s["vg"], s["id"])
        i_on = float(np.abs(if_).max())
        cycles.append(s["cycles"])
        sheets.append(s["sheet"])
        vth_f.append(vth_constant_current(vf, if_, ic))
        vth_r.append(vth_constant_current(vr, ir, ic))
        ion.append(i_on)
        ss.append(subthreshold_swing(vf, if_, i_on))
    out = {k: np.asarray(v, float) for k, v in
           (("cycles", cycles), ("vth_fwd", vth_f), ("vth_rev", vth_r),
            ("ion", ion), ("ss", ss))}
    out["mw"] = out["vth_fwd"] - out["vth_rev"]
    out["sheets"] = sheets
    return out


# ---------------------------------------------------------------------------
# Plot style
# ---------------------------------------------------------------------------
def apply_plot_style() -> None:
    plt.rcParams.update(
        {
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
            "font.size": FONT_BASE,
            "axes.labelsize": FONT_AXIS,
            "axes.titlesize": FONT_PANEL,
            "legend.fontsize": 13,
            "xtick.labelsize": FONT_BASE,
            "ytick.labelsize": FONT_BASE,
        }
    )


def style_ax(ax: plt.Axes) -> None:
    ax.spines["top"].set_visible(True)
    ax.spines["right"].set_visible(True)
    ax.tick_params(axis="both", which="major", width=1.5, length=6)
    for label in ax.get_xticklabels() + ax.get_yticklabels():
        label.set_fontweight("bold")
    ax.grid(True, which="major", alpha=GRID_ALPHA, linewidth=0.5)


def panel_label(ax: plt.Axes, text: str, dx: float = -0.22) -> None:
    """Panel letter outside the axes. ``dx`` clears the y-label, which needs more
    room in a narrow panel than in a wide one."""
    ax.text(dx, 1.04, text, transform=ax.transAxes,
            fontsize=FONT_PANEL, fontweight="bold", va="bottom", ha="left")


def cycle_x(cycles: np.ndarray) -> np.ndarray:
    """Cycle counts placed on a log axis, with 0 (pristine) parked to the left."""
    x = np.asarray(cycles, float).copy()
    x[x <= 0] = PRISTINE_X
    return x


def style_cycle_axis(ax: plt.Axes, cycles: np.ndarray) -> None:
    """Log cycle axis whose leftmost tick is the pristine reference."""
    ax.set_xscale("log")
    pos = np.asarray(cycles, float)
    pos = pos[pos > 0]
    n_dec = int(np.ceil(np.log10(pos.max()))) + 1
    # With many decades in a narrow panel, label every other one.
    step = 2 if n_dec > 6 else 1
    exps = list(range(0, n_dec, step))
    if exps[-1] != n_dec - 1:
        exps.append(n_dec - 1)
    ticks = [PRISTINE_X] + [10.0 ** e for e in exps]
    ax.set_xticks(ticks)
    ax.set_xticklabels(["pristine"] + [f"$10^{{{e}}}$" for e in exps])
    ax.set_xlim(PRISTINE_X * 0.5, 10.0 ** (n_dec - 1) * 2.2)
    ax.set_xlabel("Endurance cycles", fontweight="bold")
    ax.tick_params(axis="x", which="minor", length=0)
    for lab in ax.get_xticklabels():
        lab.set_fontsize(13)


def _plot_vs_cycles(ax: plt.Axes, cycles: np.ndarray, y: np.ndarray, color: str,
                    label: str, marker: str = "o") -> None:
    """One metric against cycles: pristine as an open reference, then the sweep."""
    x = cycle_x(cycles)
    m = np.isfinite(y)
    ax.plot(x[m], y[m], marker + "-", color=color, lw=LW_CURVE, ms=MARKERSIZE,
            mfc="white", mew=1.6, label=label)
    pr = m & (np.asarray(cycles) <= 0)
    if pr.any():
        ax.plot(x[pr], y[pr], marker, color=color, ms=MARKERSIZE + 3, mfc=color)


# ---------------------------------------------------------------------------
# Panels a / b / c -- shared by the overview and the standalone exports.
# A standalone export carries no panel letter: it is a figure, not a panel.
# ---------------------------------------------------------------------------
def cycle_norm(cycles: np.ndarray) -> Normalize:
    """Colour scale over log10(cycles), excluding the pristine reference."""
    cycled = np.asarray(cycles, float)
    cycled = cycled[cycled > 0]
    return Normalize(vmin=0.0,
                     vmax=float(np.log10(cycled.max())) if cycled.size else 1.0)


def add_cycle_colorbar(fig: plt.Figure, ax: plt.Axes, norm: Normalize) -> None:
    sm = plt.cm.ScalarMappable(norm=norm, cmap=plt.get_cmap(CMAP_CYCLES))
    cbar = fig.colorbar(sm, ax=ax, pad=0.02)
    cbar.set_label("Endurance cycles", fontweight="bold", fontsize=14)
    ticks = list(range(0, int(norm.vmax) + 1))
    cbar.set_ticks(ticks)
    cbar.set_ticklabels([f"$10^{{{e}}}$" for e in ticks])
    cbar.ax.tick_params(labelsize=12)


def draw_transfer(ax: plt.Axes, sweeps: list, norm: Normalize,
                  panel: Optional[str] = None) -> None:
    """Id-Vg dual sweeps, one trace pair per cycled checkpoint, shaded by count.

    The pristine sweep is not drawn here -- the panel is the cycle ramp and its
    colorbar, nothing else. It still sets the reference in every other panel.
    """
    cmap = plt.get_cmap(CMAP_CYCLES)
    for s in sweeps:
        if s["cycles"] <= 0:
            continue
        (vf, if_), (vr, ir) = split_branches(s["vg"], s["id"])
        color = cmap(norm(np.log10(s["cycles"])))
        ax.semilogy(vf, np.abs(if_), ls="-", color=color, lw=1.6, zorder=3)
        ax.semilogy(vr, np.abs(ir), ls="-", color=color, lw=1.6, zorder=3)
    ax.axhline(IC_A, color="k", ls=":", lw=1.4)
    ax.set_xlabel("Gate voltage $V_G$ (V)", fontweight="bold")
    ax.set_ylabel("$|I_D|$ (A)", fontweight="bold")
    ax.set_ylim(1e-14, None)
    style_ax(ax)
    if panel:
        panel_label(ax, panel, dx=-0.08)


def draw_vth(ax: plt.Axes, met: dict, panel: Optional[str] = None) -> None:
    """Forward and reverse threshold voltage against cycles."""
    cycles = met["cycles"]
    _plot_vs_cycles(ax, cycles, met["vth_fwd"], C_FWD, "forward ($-$ to $+$)", "o")
    _plot_vs_cycles(ax, cycles, met["vth_rev"], C_REV, "reverse ($+$ to $-$)", "s")
    ax.set_ylabel("$V_{th}$ (V)", fontweight="bold")
    style_cycle_axis(ax, cycles)
    style_ax(ax)
    if panel:
        panel_label(ax, panel)
    ax.legend(loc="center left", frameon=True, framealpha=0.9)


def draw_memory_window(ax: plt.Axes, met: dict, panel: Optional[str] = None) -> None:
    """Memory window against cycles, with the pristine window as the reference."""
    cycles, mw = met["cycles"], met["mw"]
    _plot_vs_cycles(ax, cycles, mw, C_MW, "memory window", "D")
    pr = np.asarray(cycles) <= 0
    if pr.any() and np.isfinite(mw[pr][0]):
        mw0 = float(mw[pr][0])
        ax.axhline(mw0, color="k", ls="--", lw=1.4, label=f"pristine: {mw0:.2f} V")
    ax.set_ylabel("Memory window (V)", fontweight="bold")
    style_cycle_axis(ax, cycles)
    style_ax(ax)
    if panel:
        panel_label(ax, panel)
    ax.legend(loc="lower left", frameon=True, framealpha=0.9)


def _save_single(fig: plt.Figure, out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=DPI, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    print(f"[ok] wrote {out_path}")


def plot_panels_separately(sweeps: list, met: dict, outdir: Path, stem: str) -> None:
    """Panels a, b and c as standalone figures, without the panel letters."""
    norm = cycle_norm(met["cycles"])

    fig, ax = plt.subplots(figsize=FIG_TRANSFER_SIZE)
    draw_transfer(ax, sweeps, norm)
    add_cycle_colorbar(fig, ax, norm)
    _save_single(fig, outdir / f"fefet_endurance_{stem}_transfer.png")

    fig, ax = plt.subplots(figsize=FIG_PANEL_SIZE)
    draw_vth(ax, met)
    fig.tight_layout()
    _save_single(fig, outdir / f"fefet_endurance_{stem}_vth.png")

    fig, ax = plt.subplots(figsize=FIG_PANEL_SIZE)
    draw_memory_window(ax, met)
    fig.tight_layout()
    _save_single(fig, outdir / f"fefet_endurance_{stem}_window.png")


# ---------------------------------------------------------------------------
# Figure: one workbook
# ---------------------------------------------------------------------------
def plot_workbook(sweeps: list, met: dict, out_path: Path, title: str) -> None:
    fig = plt.figure(figsize=FIG_SIZE)
    gs = fig.add_gridspec(2, 3, height_ratios=(1.0, 0.85), hspace=0.45, wspace=0.36,
                          left=0.055, right=0.985, top=0.90, bottom=0.09)
    ax_a = fig.add_subplot(gs[0, 0:2])
    ax_b = fig.add_subplot(gs[0, 2])
    ax_c = fig.add_subplot(gs[1, 0])
    ax_d = fig.add_subplot(gs[1, 1])
    ax_e = fig.add_subplot(gs[1, 2])

    cycles = met["cycles"]
    norm = cycle_norm(cycles)

    draw_transfer(ax_a, sweeps, norm, panel="a")
    add_cycle_colorbar(fig, ax_a, norm)
    draw_vth(ax_b, met, panel="b")
    draw_memory_window(ax_c, met, panel="c")

    # (d) on-current ---------------------------------------------------------
    _plot_vs_cycles(ax_d, cycles, met["ion"] * 1e6, C_ION, "$I_{on}$", "^")
    ax_d.set_ylabel(r"$I_{on}$ at $V_G^{max}$ ($\mu$A)", fontweight="bold")
    ax_d.set_ylim(0, None)
    style_cycle_axis(ax_d, cycles)
    style_ax(ax_d)
    panel_label(ax_d, "d")

    # (e) subthreshold swing -------------------------------------------------
    _plot_vs_cycles(ax_e, cycles, met["ss"], C_SS, "SS", "v")
    ax_e.set_ylabel("Subthreshold swing (mV/dec)", fontweight="bold")
    style_cycle_axis(ax_e, cycles)
    style_ax(ax_e)
    panel_label(ax_e, "e")

    fig.suptitle(title, fontsize=FONT_PANEL, fontweight="bold")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=DPI, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    print(f"[ok] wrote {out_path}")


# ---------------------------------------------------------------------------
# Figure: comparison across workbooks
# ---------------------------------------------------------------------------
def plot_comparison(runs: list, out_path: Path) -> None:
    fig, (ax_a, ax_b) = plt.subplots(1, 2, figsize=FIG_CMP_SIZE)
    all_cycles = np.concatenate([r["met"]["cycles"] for r in runs])

    for i, run in enumerate(runs):
        met = run["met"]
        color = DEVICE_COLORS[i % len(DEVICE_COLORS)]
        marker = "oDs^"[i % 4]
        _plot_vs_cycles(ax_a, met["cycles"], met["mw"], color, run["label"], marker)
        pr = np.asarray(met["cycles"]) <= 0
        mw0 = float(met["mw"][pr][0]) if pr.any() else float(met["mw"][0])
        _plot_vs_cycles(ax_b, met["cycles"], 100.0 * met["mw"] / mw0, color,
                        f"{run['label']} ({mw0:.2f} V)", marker)

    ax_a.set_ylabel("Memory window (V)", fontweight="bold")
    style_cycle_axis(ax_a, all_cycles)
    style_ax(ax_a)
    panel_label(ax_a, "a")
    ax_a.legend(loc="lower left", frameon=True, framealpha=0.9)

    ax_b.axhline(100.0, color="k", ls="--", lw=1.4)
    ax_b.set_ylabel("Memory window (% of pristine)", fontweight="bold")
    style_cycle_axis(ax_b, all_cycles)
    style_ax(ax_b)
    panel_label(ax_b, "b")
    ax_b.legend(loc="lower left", frameon=True, framealpha=0.9)

    fig.suptitle("FeFET endurance: memory window vs cycles",
                 fontsize=FONT_PANEL, fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.94))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=DPI, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    print(f"[ok] wrote {out_path}")


# ---------------------------------------------------------------------------
# Console summary
# ---------------------------------------------------------------------------
def report(label: str, met: dict) -> None:
    print(f"\n--- {label} ---")
    print(f"{'sheet':>16} {'cycles':>8} {'Vth_fwd':>9} {'Vth_rev':>9} "
          f"{'MW (V)':>8} {'MW (%)':>8} {'Ion (uA)':>9} {'SS':>8}")
    cycles, mw = met["cycles"], met["mw"]
    pr = np.asarray(cycles) <= 0
    mw0 = float(mw[pr][0]) if pr.any() and np.isfinite(mw[pr][0]) else float(mw[0])
    for i, sheet in enumerate(met["sheets"]):
        c = "pristine" if cycles[i] <= 0 else f"1e{int(round(np.log10(cycles[i])))}"
        print(f"{sheet:>16} {c:>8} {met['vth_fwd'][i]:+9.3f} {met['vth_rev'][i]:+9.3f} "
              f"{mw[i]:8.3f} {100.0 * mw[i] / mw0:8.1f} {met['ion'][i] * 1e6:9.2f} "
              f"{met['ss'][i]:8.1f}")

    last = len(cycles) - 1
    print(f"window {mw0:.3f} V -> {mw[last]:.3f} V "
          f"({100.0 * mw[last] / mw0:.1f} % after {cycles[last]:.0e} cycles, "
          f"{1e3 * (mw[last] - mw0):+.0f} mV)")
    d_f = met["vth_fwd"][last] - met["vth_fwd"][0]
    d_r = met["vth_rev"][last] - met["vth_rev"][0]
    print(f"  of which forward {1e3 * d_f:+.0f} mV, reverse {1e3 * d_r:+.0f} mV "
          f"-- the larger one carries the loss")

    # A step in Ion between neighbouring checkpoints is a measurement break
    # (re-probe, contact), not endurance; say so rather than plotting past it.
    ion = met["ion"]
    for i in range(1, len(ion)):
        if ion[i - 1] > 0 and abs(ion[i] - ion[i - 1]) / ion[i - 1] > ION_STEP_WARN:
            print(f"  [check] Ion steps {ion[i - 1] * 1e6:.2f} -> {ion[i] * 1e6:.2f} uA "
                  f"between {met['sheets'][i - 1]!r} and {met['sheets'][i]!r}; "
                  "likely a re-contact between runs, not device degradation")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main(argv: Optional[list] = None) -> int:
    ap = argparse.ArgumentParser(description="Plot FeFET endurance workbooks.")
    ap.add_argument("files", nargs="*", type=Path,
                    help=f"workbooks (default: {' '.join(DEFAULT_GLOBS)} in {ROOT})")
    ap.add_argument("--outdir", type=Path, default=OUTDIR)
    ap.add_argument("--ic", type=float, default=IC_A,
                    help=f"constant-current Vth criterion in A (default {IC_A:g})")
    ap.add_argument("--no-comparison", action="store_true")
    ap.add_argument("--no-panels", action="store_true",
                    help="skip the standalone exports of panels a, b and c")
    args = ap.parse_args(argv)

    paths = list(args.files)
    if not paths:
        for pattern in DEFAULT_GLOBS:
            paths.extend(sorted(ROOT.glob(pattern)))
    paths = [p for p in paths if p.is_file()]
    if not paths:
        print(f"[error] no workbooks found ({', '.join(DEFAULT_GLOBS)}) in {ROOT}",
              file=sys.stderr)
        return 1

    apply_plot_style()
    runs = []
    for path in paths:
        sweeps = read_sweeps(path)
        if not sweeps:
            print(f"[warn] {path.name}: no Id-Vg sheets; skipped")
            continue
        mapping = ", ".join(
            f"{s['sheet']}->" + ("pristine" if s["cycles"] <= 0
                                 else f"1e{int(round(np.log10(s['cycles'])))}")
            for s in sweeps)
        print(f"[in ] {path.name}: {len(sweeps)} checkpoints  [{mapping}]")
        met = metrics(sweeps, ic=args.ic)
        label = path.stem
        stem = re.sub(r"[^A-Za-z0-9]+", "_", label).strip("_")
        plot_workbook(sweeps, met, args.outdir / f"fefet_endurance_{stem}.png",
                      f"FeFET endurance -- {label}  ($V_D$ = 0.1 V)")
        if not args.no_panels:
            plot_panels_separately(sweeps, met, args.outdir, stem)
        runs.append({"label": label, "met": met})

    if len(runs) > 1 and not args.no_comparison:
        plot_comparison(runs, args.outdir / "fefet_endurance_comparison.png")

    for run in runs:
        report(run["label"], run["met"])
    print()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
